"""Load `plm_item` records into Postgres.

Primary source: the Centric 8 PLM API (`style_to_plm_item` + `refresh_plm_items`, the live
sync source). The Excel `Wrike` sheet / synthetic-CSV readers below are retained as test and
local-seed scaffolding. Both paths feed the same `load_plm_items` upsert.
"""
from __future__ import annotations

import csv
import html
import re
from datetime import datetime, timezone
from pathlib import Path

from description import build_description
from state import EPOCH


def _read_excel_wrike_sheet(path) -> list[dict]:
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        # openpyxl warns about an unsupported data-validation extension in this file; benign.
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Wrike"]
        rows = ws.iter_rows(values_only=True)
        header = ["" if h is None else str(h) for h in next(rows)]
        out = []
        for r in rows:
            if all(c is None or str(c).strip() == "" for c in r):
                continue  # skip blank rows
            out.append({header[i]: ("" if c is None else str(c)) for i, c in enumerate(r)})
    return out


def _read_csv(path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def read_source_rows(excel_path, csv_path=None) -> list[dict]:
    """Return the union of the Excel `Wrike` sheet and (optional) synthetic CSV rows."""
    rows = _read_excel_wrike_sheet(excel_path)
    if csv_path and Path(csv_path).exists():
        rows += _read_csv(csv_path)
    return rows

# Source column header -> plm_item field name.
COLUMN_MAP = {
    "Key": "plm_internal_id",
    "PLM - Item #": "item_number",
    "PLM Item Name": "item_name",
    "Title": "title",
    "Folder": "folder",
    "Workflow": "workflow",
    "Status": "status",
    "Custom status": "custom_status",
    "Priority": "priority",
    "End Date": "end_date",
    "Description": "description",
    "Print Method": "print_method",
    "Previous Item No": "previous_item_no",
    "PLM - Contents": "contents",
    "PLM - Material Codes": "material_codes",
    "Brand Category": "brand_category",
    "Customer": "customer",
    "Product Category": "product_category",
    "Brand": "brand",
    "Season": "season",
    "PLM - Design Request": "design_request",
    "Design Brief / Additional Notes": "design_brief",
    "Image Link": "image_link",
}


PLM_ITEM_COLUMNS = [
    "plm_internal_id", "item_number", "family_id", "prefix", "code", "item_name",
    "customer", "title", "folder", "workflow", "status", "custom_status", "priority",
    "end_date", "description", "print_method", "previous_item_no", "contents",
    "material_codes", "brand_category", "product_category", "brand", "season",
    "design_request", "design_brief", "image_link", "ready_for_wrike",
    "created_at", "modified_at",
]


def load_plm_items(conn, items: list[dict]) -> None:
    """Upsert plm_item rows (keyed on plm_internal_id). Idempotent."""
    cols = ", ".join(PLM_ITEM_COLUMNS)
    placeholders = ", ".join(f"%({c})s" for c in PLM_ITEM_COLUMNS)
    updates = ", ".join(
        f"{c} = EXCLUDED.{c}" for c in PLM_ITEM_COLUMNS if c != "plm_internal_id"
    )
    sql = (
        f"INSERT INTO plm_item ({cols}) VALUES ({placeholders}) "
        f"ON CONFLICT (plm_internal_id) DO UPDATE SET {updates}"
    )
    with conn.cursor() as cur:
        cur.executemany(sql, items)
    conn.commit()


def item_prefix(item_number: str) -> str:
    """Leading alpha prefix of an item number, e.g. 'WP-71511-006-319' -> 'WP'."""
    match = re.match(r"^[A-Za-z]+", item_number)
    return match.group(0) if match else ""


def to_plm_item(row: dict, *, now) -> dict:
    item = {field: row.get(col) for col, field in COLUMN_MAP.items()}

    item_number = item["item_number"]
    parts = item_number.split("-")
    item["family_id"] = item_number[:8]
    item["prefix"] = item_prefix(item_number)
    item["code"] = parts[1] if len(parts) > 1 else ""

    # synthesized control columns (absent from the source Excel):
    item["ready_for_wrike"] = True
    item["created_at"] = now
    item["modified_at"] = now
    return item


# ---------------------------------------------------------------------------
# Centric 8 API source: map styles -> plm_item and refresh the local mirror.
# (The Excel/CSV path above is retained as test/local-seed scaffolding.)
# ---------------------------------------------------------------------------


def strip_html(text: str | None) -> str:
    """Plain text from a Centric HTML field: drop tags, unescape entities, trim blank
    lines (e.g. mgf_test_material -> contents)."""
    if not text:
        return ""
    text = re.sub(r"</(p|div)>|<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", "", text)
    text = html.unescape(text)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _parse_centric_timestamp(value) -> datetime | None:
    """Parse a Centric `_modified_at` value to an aware UTC datetime; None if absent.

    The sandbox's exact format is unconfirmed (see project_docs/centric_8_api.md), so
    accept ISO 8601 and Centric's 'yyyy/mm/ddThh:mm:ss', assuming UTC when naive. A
    present-but-unparseable value raises loudly rather than silently defaulting.
    """
    if not value:
        return None
    text = str(value).strip().replace("Z", "+00:00")
    parsers = (
        datetime.fromisoformat,
        lambda t: datetime.strptime(t, "%Y/%m/%dT%H:%M:%S"),
        lambda t: datetime.strptime(t, "%Y/%m/%d %H:%M:%S"),
    )
    for parse in parsers:
        try:
            dt = parse(text)
        except ValueError:
            continue
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    raise ValueError(f"unrecognized Centric _modified_at timestamp: {value!r}")


def _format_centric_timestamp(dt: datetime) -> str:
    """A datetime as Centric's `modified_after` query string (yyyy/mm/ddThh:mm:ss, UTC)."""
    return dt.astimezone(timezone.utc).strftime("%Y/%m/%dT%H:%M:%S")


def style_to_plm_item(style: dict, resolve, *, now: datetime) -> dict:
    """Map one Centric style (raw API dict) to a plm_item record - see specs/002 §3.

    `resolve(endpoint, ref_id)` turns a reference id into its display name
    (CentricClient.resolve_ref); `now` is the modified_at fallback when a style carries no
    `_modified_at`. created_at reuses modified_at - the API exposes no creation date.
    """
    item_number = style.get("mgf_item_identifier") or ""
    parts = item_number.split("-")
    material_codes = style.get("mgf_material_code_item") or ""
    contents = strip_html(style.get("mgf_test_material"))
    item_name = style.get("node_name") or ""
    modified_at = _parse_centric_timestamp(style.get("_modified_at")) or now

    return {
        "plm_internal_id": style["id"],
        "item_number": item_number,
        "family_id": item_number[:8],
        "prefix": item_prefix(item_number),
        "code": parts[1] if len(parts) > 1 else "",
        "item_name": item_name,
        "customer": resolve("category2s", style.get("category_2")),          # *CORE / *CUSTOM
        "title": f"{item_number[:8]} {item_name}".strip(),
        "folder": None,
        "workflow": None,
        "status": None,
        "custom_status": None,
        "priority": None,
        "end_date": None,
        "description": build_description(material_codes, contents),
        "print_method": ", ".join(style.get("mgf_print_method") or []),
        "previous_item_no": style.get("mgf_previous_item_number") or "",
        "contents": contents,
        "material_codes": material_codes,
        "brand_category": style.get("mgf_brand_category_2") or "",
        "product_category": resolve("category1s", style.get("category_1")),  # drives author map
        "brand": resolve("collections", style.get("collection")),
        "season": "",                       # /seasons broken on the sandbox -> blank
        "design_request": style.get("mgf_item_description") or "",
        "design_brief": "",
        "image_link": style.get("mgf_image_link") or "",
        "ready_for_wrike": bool(style.get("mgf_ready_for_wrike")),
        "created_at": modified_at,
        "modified_at": modified_at,
    }


def refresh_plm_items(conn, client, *, since: datetime, now: datetime) -> int:
    """Pull active styles changed since `since` from Centric, map them, and upsert into
    plm_item (the local mirror). Returns the number of rows upserted.

    The fetch is NOT gated on ready_for_wrike - reconciliation needs not-ready items too;
    the producer's delta gates on the ready_for_wrike column. `since` at/<= EPOCH (the
    first run) pulls the full active catalogue; later runs pull only the modified_after
    delta.
    """
    modified_after = None if since <= EPOCH else _format_centric_timestamp(since)
    styles = client.list_styles(modified_after=modified_after)
    items = [style_to_plm_item(s, client.resolve_ref, now=now)
             for s in styles if s.get("active")]
    if items:
        load_plm_items(conn, items)
    return len(items)
